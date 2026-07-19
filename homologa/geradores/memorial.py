"""
Gerador do Memorial Técnico Descritivo (.docx).

v0: gera o documento do zero com uma estrutura-padrão razoável, para provar o
fluxo ponta a ponta. Quando você enviar seu template real (.docx com marcadores),
trocamos esta função por preenchimento do SEU modelo — a fonte de dados (Projeto
+ ResultadoCalculo) permanece a mesma.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from ..calculos import ResultadoCalculo, calcular
from ..modelo import Projeto


def _titulo(doc: Document, texto: str, nivel: int = 1) -> None:
    h = doc.add_heading(texto, level=nivel)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _kv(doc: Document, chave: str, valor: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"{chave}: ")
    r.bold = True
    p.add_run(valor or "—")


def gerar_memorial(projeto: Projeto, destino: str | Path,
                   resultado: ResultadoCalculo | None = None) -> Path:
    if resultado is None:
        resultado = calcular(projeto)

    destino = Path(destino)
    doc = Document()

    # Estilo base
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    tit = doc.add_heading("MEMORIAL TÉCNICO DESCRITIVO", level=0)
    tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Sistema de Geração Distribuída Fotovoltaica")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. Identificação
    _titulo(doc, "1. Identificação")
    _kv(doc, "Titular", projeto.titular.nome)
    _kv(doc, "CPF/CNPJ", projeto.titular.cpf_cnpj)
    _kv(doc, "Unidade consumidora (matriz)", projeto.unidade_geradora.numero_uc)
    _kv(doc, "Endereço da instalação",
        projeto.unidade_geradora.endereco.uma_linha())
    _kv(doc, "Concessionária", projeto.concessionaria.nome)
    _kv(doc, "Classe de consumo", projeto.unidade_geradora.classe.value)
    _kv(doc, "Tipo de ligação", projeto.unidade_geradora.tipo_ligacao.value)
    _kv(doc, "Tensão de atendimento", f"{projeto.unidade_geradora.tensao_v:.0f} V")

    # 2. Descrição do sistema
    _titulo(doc, "2. Descrição do Sistema")
    doc.add_paragraph(
        f"O sistema é classificado como {resultado.classificacao.upper()} "
        f"distribuída, com potência de geração instalada de "
        f"{resultado.potencia_geracao_kwp:.2f} kWp e potência nominal de saída "
        f"de {resultado.potencia_instalada_kw:.2f} kW."
    )

    # 3. Módulos
    _titulo(doc, "3. Módulos Fotovoltaicos")
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    hdr[0].text, hdr[1].text = "Fabricante", "Modelo"
    hdr[2].text, hdr[3].text = "Pot. (Wp)", "Qtd."
    for m in projeto.gerador.modulos:
        c = tbl.add_row().cells
        c[0].text = m.fabricante
        c[1].text = m.modelo
        c[2].text = f"{m.potencia_wp:.0f}"
        c[3].text = str(m.quantidade)
    _kv(doc, "Total de módulos", str(resultado.n_modulos))

    # 4. Inversores
    _titulo(doc, "4. Inversores")
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    hdr[0].text, hdr[1].text = "Fabricante", "Modelo"
    hdr[2].text, hdr[3].text = "Pot. CA (W)", "Qtd."
    for i in projeto.gerador.inversores:
        c = tbl.add_row().cells
        c[0].text = i.fabricante
        c[1].text = i.modelo
        c[2].text = f"{i.potencia_ca_w:.0f}"
        c[3].text = str(i.quantidade)
    _kv(doc, "Total de inversores", str(resultado.n_inversores))

    # 5. Dimensionamento e proteção
    _titulo(doc, "5. Dimensionamento e Proteção")
    _kv(doc, "Relação FV/inversor", f"{resultado.relacao_fv_inversor:.2f}")
    _kv(doc, "Corrente nominal de saída CA", f"{resultado.corrente_saida_a:.2f} A")
    _kv(doc, "Disjuntor do gerador (calculado)",
        f"{resultado.disjuntor_gerador_calc_a:.2f} A")
    _kv(doc, "Disjuntor do gerador (comercial)",
        f"{resultado.disjuntor_gerador_comercial_a} A")
    _kv(doc, "Disjuntor de entrada existente",
        f"{projeto.unidade_geradora.disjuntor_entrada_a:.0f} A")

    # 6. Rateio (se houver)
    if projeto.beneficiarias:
        _titulo(doc, "6. Sistema de Compensação — Rateio")
        tbl = doc.add_table(rows=1, cols=3)
        tbl.style = "Light Grid Accent 1"
        hdr = tbl.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "UC", "Titular", "Rateio (%)"
        for b in projeto.beneficiarias:
            c = tbl.add_row().cells
            c[0].text = b.numero_uc
            c[1].text = b.titular_nome
            c[2].text = f"{b.percentual_rateio:.2f}"

    # 7. Responsável técnico
    _titulo(doc, "7. Responsável Técnico")
    rt = projeto.responsavel_tecnico
    _kv(doc, "Nome", rt.nome)
    _kv(doc, "Título", rt.titulo)
    _kv(doc, "CREA", rt.crea)
    _kv(doc, "ART", rt.art)

    doc.save(destino)
    return destino
