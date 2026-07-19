"""
Gerador da Autorização de Representação Técnica (.docx).

Documento que o cliente assina autorizando o responsável técnico a
protocolar o pedido de acesso junto à concessionária em seu nome.
Estrutura-padrão para o v0 — trocamos pelo seu texto oficial depois.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from ..modelo import Projeto, TipoPessoa


def gerar_autorizacao(projeto: Projeto, destino: str | Path) -> Path:
    destino = Path(destino)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    tit = doc.add_heading("AUTORIZAÇÃO DE REPRESENTAÇÃO TÉCNICA", level=0)
    tit.alignment = WD_ALIGN_PARAGRAPH.CENTER

    t = projeto.titular
    rt = projeto.responsavel_tecnico
    uc = projeto.unidade_geradora

    doc_titular = "CNPJ" if t.tipo_pessoa == TipoPessoa.JURIDICA else "CPF"

    corpo = doc.add_paragraph()
    corpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    corpo.add_run(
        f"Eu, {t.nome or '________________'}, portador(a) do documento "
        f"{doc_titular} nº {t.cpf_cnpj or '____________'}, titular da unidade "
        f"consumidora nº {uc.numero_uc or '____________'}, localizada em "
        f"{uc.endereco.uma_linha() or '____________'}, AUTORIZO o profissional "
        f"{rt.nome or '________________'}, {rt.titulo}, inscrito no CREA sob o "
        f"nº {rt.crea or '____________'}, a me representar tecnicamente junto à "
        f"{projeto.concessionaria.nome} em todos os atos referentes ao pedido de "
        f"acesso, conexão, vistoria e homologação do sistema de microgeração/"
        f"minigeração distribuída fotovoltaica instalado na referida unidade "
        f"consumidora."
    )

    doc.add_paragraph()
    local_data = doc.add_paragraph(
        f"{uc.endereco.cidade or '____________'} ({uc.endereco.uf}), "
        f"{projeto.data_emissao or '____ de __________ de ______'}."
    )
    local_data.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()
    doc.add_paragraph()
    ass = doc.add_paragraph("_" * 45)
    ass.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nome = doc.add_paragraph(t.nome or "Titular da unidade consumidora")
    nome.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cpf = doc.add_paragraph(f"{doc_titular}: {t.cpf_cnpj or ''}")
    cpf.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(destino)
    return destino
